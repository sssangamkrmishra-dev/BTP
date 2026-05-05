#!/usr/bin/env python3
"""
Comprehensive Markdown-to-DOCX converter for BTP design documents.

Uses python-docx for full formatting control, Mermaid CLI for diagram
rendering, and content-based matching to correctly identify every
ASCII art block.
"""

import os
import re
import subprocess
import json
import textwrap
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

DOCS_DIR = os.path.dirname(os.path.abspath(__file__))
DIAGRAMS_DIR = os.path.join(DOCS_DIR, "diagrams")
os.makedirs(DIAGRAMS_DIR, exist_ok=True)

# ===================================================================
# MERMAID DIAGRAM DEFINITIONS
# Each entry: (unique_content_match_string, mermaid_code, width_inches)
# The match string is searched in the ASCII art block body.
# ===================================================================

II_DIAGRAM_MAP = [
    # 3.1 Full System Architecture
    ("LAYER 1: DRONE / RPA PLATFORM", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}}}%%
flowchart LR
    subgraph L1["LAYER 1<br/>DRONE / RPA"]
        DRONE["Drone / RPA<br/>video, images,<br/>telemetry, metadata"]
    end

    subgraph L2["LAYER 2 : EDGE MALWARE DETECTION ENGINE"]
        direction TB
        subgraph ROW1[" "]
            direction LR
            M1["<b>M1: Ingestion<br/>Interceptor</b>"]:::thismod
            M2["M2: Metadata<br/>Extractor"]
            M3["M3: Game-Theoretic<br/>Threat Estimator"]
            M4["M4: Inspection<br/>Strategy Selector"]
            M1 -->|"IngestResult"| M2 --> M3 -->|"T_S"| M4
        end
        subgraph ROW2[" "]
            direction LR
            M5["M5: Multi-Layer<br/>Malware Detection"]
            M6["M6: Metadata<br/>Sanitizer"]
            M7["M7: Threat Intel<br/>Correlator"]
            M8["M8: Response &amp;<br/>Quarantine Mgr"]
            M5 --> M6 --> M7 --> M8
        end
        ROW1 -->|"inspection level"| ROW2
    end

    subgraph L3["LAYER 3<br/>LOGGING &amp;<br/>FEEDBACK"]
        direction TB
        M9["M9: Security<br/>Dashboard"]
        GCC["<b>Ground Control<br/>Center</b>"]:::gcc
        M9 <-->|"packets<br/>(MQTT / gRPC<br/>over TLS)"| GCC
    end

    DRONE --> M1
    M8 --> M9
    M9 -.->|"Uplink cmds"| M1

    classDef thismod fill:#BBDEFB,stroke:#0D47A1,stroke-width:3px
    classDef gcc fill:#E8F5E9,stroke:#388E3C,stroke-width:2px
    style L1 fill:#E8F5E9,stroke:#2E7D32,stroke-width:2px
    style L2 fill:#FFF8E1,stroke:#F57F17,stroke-width:2px
    style L3 fill:#F3E5F5,stroke:#7B1FA2,stroke-width:2px
    style ROW1 fill:none,stroke:none
    style ROW2 fill:none,stroke:none
""", 6.5),

    # 3.2 Boundary Context
    ("TRUST BOUNDARY", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}}}%%
flowchart TB
    DRONE["<b>Drone / RPA Platform</b>"]
    GS["<b>Ground Station<br/>Adapter</b><br/><i>Parses native protocol<br/>(MAVLink, DJI MSDK, etc.)<br/>and builds submission dict</i>"]:::adapter

    subgraph TB["TRUST BOUNDARY"]
        direction TB
        subgraph II["<b>INGESTION INTERCEPTOR</b>"]
            direction TB
            S0["<b>Stage 0:</b> Packet Receiver<br/>(HMAC verify, reassemble,<br/>decode binary TLV)"]:::stage0
            S17["<b>Stages 1-7:</b> Validate &rarr; Authenticate &rarr;<br/>Extract metadata &rarr; Analyse &rarr;<br/>Checksum &rarr; Catalog &rarr; Assemble"]
            S0 --> S17
        end
        GTE["<b>Game-Theoretic Threat Estimator</b><br/>Consumes ingest_metadata + artifact_records<br/>Produces T_S (0.0 - 1.0)"]
        UR["<b>Uplink Receiver</b><br/>QUARANTINE, RELEASE, REVOKE_DEVICE,<br/>UPDATE_ZONE_RISK, UPDATE_CONFIG, FORCE_RESCAN"]
        DASH["<b>Security Dashboard</b><br/>Visualization, audit logs,<br/>alerting, uplink dispatch"]
        II -->|"IngestResult or error report"| GTE
        DASH -->|"commands"| UR
    end

    GCC["<b>Ground Control<br/>Center</b>"]:::gcc

    DRONE -->|"UDP packets<br/>(Stage 0 — custom RPAD)"| II
    DRONE -->|"native protocol"| GS
    GS -->|"dict via process()<br/>(primary path for<br/>real drones)"| II
    DASH <-->|"packets<br/>(MQTT / gRPC over TLS)"| GCC

    classDef stage0 fill:#FFF9C4,stroke:#F9A825,stroke-width:2px
    classDef adapter fill:#E8F5E9,stroke:#388E3C,stroke-width:2px
    classDef gcc fill:#E8F5E9,stroke:#388E3C,stroke-width:2px
    style TB fill:#FFF8E1,stroke:#F57F17,stroke-width:2px
    style II fill:#E3F2FD,stroke:#1565C0,stroke-width:2px
""", 6.5),

    # 4.2 Component Dependency Diagram
    ("(IngestionInterceptor class)", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}}}%%
flowchart TB
    INT["<b>interceptor.py</b><br/>IngestionInterceptor<br/>8-stage pipeline"]:::main

    PR["<b>packet_receiver.py</b><br/>Stage 0: UDP listener<br/>+ ReassemblyBuffer + GC"]:::stage0
    PROTO["<b>protocol/</b><br/>packet codec +<br/>submission_codec (TLV)"]:::stage0

    CFG["config.py"]
    VAL["validator.py"]
    AUTH["authenticator.py"]
    META["metadata_extractor.py"]
    PA["payload_analyzer.py"]
    CV["checksum_verifier.py"]
    AM["artifact_manager.py"]
    UP["uplink.py"]
    MOD["models.py"]

    PR -.->|"on_submission()"| INT
    PR --> PROTO
    INT --> VAL & AUTH & META & PA & CV & AM
    CFG -.->|"shared config"| VAL & AUTH & META & PA & CV & AM & PR
    UP -.->|"Modifies registry"| AUTH
    MOD -.->|"imported by all"| INT

    classDef main fill:#E3F2FD,stroke:#1565C0,stroke-width:3px
    classDef stage0 fill:#FFF9C4,stroke:#F9A825,stroke-width:2px
""", 6.0),

    # 4.3 Data Flow Through Pipeline
    ("validate_submission()", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '15px'}}}%%
flowchart TB
    UDP(["UDP packets"])
    subgraph S0["Stage 0 (optional, when start_packet_listener is running)"]
        direction LR
        PR["PacketReceiver<br/>HMAC verify + reassemble"] --> DEC["decode_submission()"]
    end

    IN(["Raw dict"])
    VS["validate_submission()"]
    REJ1["REJECT<br/>(success=False)"]:::reject
    DS["DroneSubmission.from_dict()"]
    AA["Authenticator.authenticate()"]
    REJ2["REJECT"]:::reject
    EX["extract_mission_context()<br/>extract_geo_metadata()<br/>extract_telemetry_summary()<br/>extract_additional_metadata()"]
    subgraph LOOP["FOR EACH payload"]
        AP["analyze_payload()"] -->|"security_flags[]"| VC["verify_checksum()"]
        VC -->|"checksum_verified"| CR["create_artifact_record()"]
    end
    BM["Build IngestMetadata<br/>(aggregate flags, zone risk, notes)"]
    OUT(["IngestResult(success=True)"]):::success

    UDP --> S0 --> IN
    IN --> VS
    VS -->|"errors?"| REJ1
    VS -->|"OK"| DS --> AA
    AA -->|"rejected?"| REJ2
    AA -->|"OK"| EX --> LOOP --> BM --> OUT

    classDef reject fill:#FFCDD2,stroke:#C62828
    classDef success fill:#C8E6C9,stroke:#2E7D32
    style S0 fill:#FFF9C4,stroke:#F9A825,stroke-width:2px
""", 5.5),

    # 5.4 Authenticator Architecture
    ("(Ed25519 planned)", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '17px'}}}%%
flowchart TB
    subgraph AUTH["<b>Authenticator</b><br/>(orchestration layer)"]
        direction LR
        subgraph DR["<b>DeviceRegistry</b>"]
            direction TB
            L["lookup(drone_id)"]
            R["register_device()"]
            RV["revoke_device()"]
            LD["list_devices()"]
            BK["<i>Backends:</i><br/>In-memory dict<br/>JSON file"]
        end
        subgraph SV["<b>SignatureVerifier</b>"]
            direction TB
            VER["verify(drone_id, signature,<br/>payload_hash)"]
            SCH["Schemes: HMAC-SHA256<br/>(Ed25519 planned)"]
            KS["Key store:<br/>drone_id &rarr; secret"]
        end
    end

    style AUTH fill:#E3F2FD,stroke:#1565C0,stroke-width:2px
    style DR fill:#E8F5E9,stroke:#388E3C,stroke-width:2px
    style SV fill:#FFF9C4,stroke:#F9A825,stroke-width:2px
""", 6.5),

    # 5.9 Uplink Architecture
    ("memory (unit testing)", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}}}%%
flowchart LR
    CC["<b>Control Center /<br/>Security Dashboard</b><br/><br/>Dispatches:<br/>Quarantine · Release<br/>Revoke Device<br/>Update Zone Risk<br/>Update Config<br/>Force Rescan"]
    UR["<b>UplinkReceiver</b><br/><br/><i>Modes:</i><br/>memory (unit testing)<br/>file (JSON file watch)<br/>gRPC (production)<br/>MQTT (production)<br/><br/>Queue + acknowledgement"]
    UCH["<b>UplinkCommandHandler</b><br/><br/><i>Dispatch table:</i><br/>QUARANTINE &rarr; add to set<br/>RELEASE &rarr; remove<br/>REVOKE &rarr; registry mut<br/>ZONE_RISK &rarr; risk map mut"]

    CC -->|"commands"| UR --> UCH

    style CC fill:#E3F2FD,stroke:#1565C0,stroke-width:2px
    style UR fill:#FFF9C4,stroke:#F9A825,stroke-width:2px
    style UCH fill:#E8F5E9,stroke:#388E3C,stroke-width:2px
""", 6.5),

    # 5.10 packet_receiver.py - Three classes
    ("PacketReceiverStats", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '14px'}}}%%
flowchart TB
    subgraph RB["<b>ReassemblyBuffer</b>"]
        direction TB
        RB1["<i>Per-session reassembly state</i><br/>Keyed by (drone_id, session_id)"]
        RB2["chunks: dict[seq &rarr; bytes]<br/>total_chunks: int (set in SESSION_START)<br/>started_at: float (for timeout)"]
        RB3["add_chunk(seq, payload) &rarr; bool<br/>is_complete() &rarr; bool<br/>is_expired(timeout) &rarr; bool<br/>assemble() &rarr; bytes"]
        RB1 --- RB2 --- RB3
    end

    subgraph PR["<b>PacketReceiver</b>"]
        direction TB
        PR1["<i>UDP socket listener</i><br/>(daemon thread)"]
        PR2["start() / stop()<br/>bound_address (property)<br/>update_key_store()<br/>on_submission callback"]
        PR3["<i>Internals:</i><br/>_listen_loop · _handle_packet<br/>_handle_session_start / _chunk / _session_end<br/>_finalize_session · _gc_expired_sessions"]
        PR1 --- PR2 --- PR3
    end

    subgraph PRS["<b>PacketReceiverStats</b>"]
        direction TB
        PRS1["<i>Counters surfaced via</i><br/><i>interceptor.packet_stats</i>"]
        PRS2["packets_received<br/>packets_dropped_malformed<br/>packets_dropped_hmac<br/>packets_dropped_unknown_drone<br/>packets_dropped_replay<br/>packets_dropped_orphaned<br/>packets_dropped_session_overflow"]
        PRS3["sessions_started · sessions_completed<br/>sessions_expired · sessions_truncated<br/>submissions_decoded<br/>submissions_decode_failed"]
        PRS1 --- PRS2 --- PRS3
    end

    PR -.->|"creates per session"| RB
    PR -.->|"updates"| PRS

    style RB fill:#E8F5E9,stroke:#388E3C,stroke-width:2px
    style PR fill:#E3F2FD,stroke:#1565C0,stroke-width:2px
    style PRS fill:#FFF9C4,stroke:#F9A825,stroke-width:2px
""", 6.5),

    # 8.1 Main Processing Flow
    ("Caller                Interceptor          Validator       Authenticator", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}, 'sequence': {'mirrorActors': false, 'width': 200}}}%%
sequenceDiagram
    participant C as Caller
    participant I as Interceptor
    participant V as Validator
    participant A as Authenticator
    participant ME as MetadataExtractor
    participant PA as PayloadAnalyzer
    participant CV as ChecksumVerifier

    C->>I: process(drone_json)
    I->>I: _process_uplink()
    I->>V: validate_submission()
    V-->>I: (errors=[], warnings=[])
    I->>I: DroneSubmission.from_dict()
    I->>I: compute_bytes_checksum()
    I->>A: authenticate(drone_id, signature, hash)
    A-->>I: AuthResult(status="authenticated")
    I->>ME: extract_mission_context()
    ME-->>I: mission context
    I->>ME: extract_geo_metadata()
    ME-->>I: validated geo
    I->>ME: extract_telemetry_summary()
    ME-->>I: telemetry + anomalies
    I->>ME: extract_additional_metadata()
    ME-->>I: sanitized metadata
    rect rgb(240,248,255)
    loop FOR EACH payload
        I->>PA: analyze_payload()
        PA-->>I: security_flags[]
        I->>CV: verify_checksum()
        CV-->>I: True / False / None
        I->>I: create_artifact_record()
    end
    end
    I->>I: Build IngestMetadata
    I-->>C: IngestResult(success=True)
""", 6.5),

    # 8.2 Authentication Flow
    ("Interceptor           Authenticator        DeviceRegistry    SignatureVerifier", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}, 'sequence': {'mirrorActors': false}}}%%
sequenceDiagram
    participant I as Interceptor
    participant A as Authenticator
    participant DR as DeviceRegistry
    participant SV as SignatureVerifier

    I->>A: authenticate(drone_id, signature, payload_hash)
    A->>DR: lookup(drone_id)

    alt CASE 1: Not Found
        DR-->>A: None
        alt policy == "reject"
            A-->>I: AuthResult(status="rejected")
        else policy == "flag" / "allow"
            A-->>I: AuthResult(status="unknown")
        end
    else CASE 2: Found but Revoked
        DR-->>A: {revoked: True}
        A-->>I: AuthResult(status="rejected")
    else CASE 3: Found and Active
        DR-->>A: {trusted, reputation}
        opt signature provided
            A->>SV: verify(drone_id, signature, hash)
            SV-->>A: {valid, reason}
            alt valid
                Note over A: status = "authenticated"
            else not valid
                Note over A: trusted=False, status="untrusted"
            end
        end
        A-->>I: AuthResult(status, reputation, trusted)
    end
""", 6.5),

    # 8.3 Uplink Command Flow
    ("Control Center        UplinkReceiver     UplinkCommandHandler", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}, 'sequence': {'mirrorActors': false}}}%%
sequenceDiagram
    participant CC as Control Center
    participant UR as UplinkReceiver
    participant UCH as CommandHandler
    participant A as Authenticator
    participant I as Interceptor

    CC->>UR: push_command(REVOKE_DEVICE, "DRN-005")
    Note over UR: command queued
    Note over I,A: Later, during process() call
    I->>UR: poll_commands()
    UR-->>I: [UplinkCommand(...)]
    I->>UCH: handle(command)
    UCH->>A: revoke_device("DRN-005")
    A-->>UCH: {trusted: False, revoked: True}
    UCH-->>I: {"status": "revoked"}
    I->>UR: acknowledge(cmd_id)
    Note over UR: command marked processed
""", 6.0),

    # 8.4 Stage 0 Packet Reception & Reassembly Flow
    ("Drone Platform        PacketReceiver        ReassemblyBuffer    process()", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '15px'}, 'sequence': {'mirrorActors': false, 'width': 220}}}%%
sequenceDiagram
    participant D as Drone Platform
    participant PR as PacketReceiver
    participant RB as ReassemblyBuffer
    participant P as process()

    Note over D: encode_submission then fragment into N chunks

    D->>PR: SESSION_START (session_id, total=N)
    activate PR
    PR->>PR: decode_packet_unverified()
    PR->>PR: lookup HMAC key by drone_id
    PR->>PR: verify_hmac()
    PR->>RB: create ReassemblyBuffer
    activate RB

    D->>PR: CHUNK seq=0
    PR->>PR: decode and verify HMAC
    PR->>RB: add_chunk(0, payload)

    D->>PR: CHUNK seq=1
    PR->>PR: decode and verify HMAC
    PR->>RB: add_chunk(1, payload)

    Note over D,RB: ... more CHUNKs ...

    D->>PR: CHUNK seq=N-1
    PR->>RB: add_chunk(N-1, payload)
    RB-->>PR: is_complete returns True

    rect rgb(232, 245, 233)
    Note over PR,P: EAGER FINALIZATION
    PR->>RB: pop session
    deactivate RB
    PR->>PR: assemble() to bytes
    PR->>PR: decode_submission() to dict
    PR->>P: on_submission(dict)
    activate P
    Note over P: process() runs Stages 1 through 7
    deactivate P
    end

    D->>PR: SESSION_END
    Note over PR: session_id unknown - benign late END (no error, no count)
    deactivate PR
""", 6.5),

    # 8.4 Rejection / truncation paths
    ("packets_dropped_orphaned += 1", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '15px'}, 'sequence': {'mirrorActors': false}}}%%
sequenceDiagram
    participant D as Drone Platform
    participant PR as PacketReceiver

    Note over D: SESSION_START dropped on the wire
    D->>PR: CHUNK seq=k
    Note over PR: unknown session_id then packets_dropped_orphaned plus one

    Note over D: Tampered packet
    D->>PR: Tampered CHUNK
    Note over PR: HMAC verification failed then packets_dropped_hmac plus one

    Note over D: Pcap replay
    D->>PR: Replayed CHUNK seq=k
    Note over PR: add_chunk returns False then packets_dropped_replay plus one

    Note over D: Truncation attack or loss
    D->>PR: CHUNKs 0..k-1 sent
    Note over D: CHUNK k missing
    D->>PR: SESSION_END
    Note over PR: assemble() raises then sessions_truncated plus one and incomplete_packet_stream log

    Note over D: Half-open session flooding
    D->>PR: SESSION_START opens session
    Note over D: no further packets for 30s
    Note over PR: GC pass detects expiry then sessions_expired plus one
""", 6.5),

    # 8.5 Rejection Flow (validation failure - was 8.4)
    ('success=False,', r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}, 'sequence': {'mirrorActors': false}}}%%
sequenceDiagram
    participant C as Caller
    participant I as Interceptor
    participant V as Validator

    C->>I: process({"drone_id": ""})
    I->>V: validate_submission()
    V-->>I: errors=["invalid_drone_id:..."]
    Note over I: stats["total_rejected"] += 1
    I-->>C: IngestResult(success=False, errors=[...])
""", 5.0),

    # 9.5 Downstream interface
    ("Stackelberg_equilibrium", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '18px'}}}%%
flowchart TB
    subgraph INPUTS["Inputs from IngestResult"]
        direction LR
        IF["insecure_flags"]
        REP["reputation R"]
        ZR["zone_risk Z"]
        FS["artifact sizes"]
        MS["mission_sensitivity"]
    end

    subgraph COMP["Computation"]
        direction LR
        IB["I_base"] --> IP["I' = I_base x sensitivity"] --> TS["T_S = Stackelberg(I', R, Z)"]
    end

    subgraph OUT["Output: T_S"]
        direction LR
        LO["<b>Low</b> &lt; 0.4<br/>Signature only"]
        ME["<b>Medium</b> 0.4-0.7<br/>Sig + AI/ML"]
        HI["<b>High</b> &ge; 0.7<br/>Full inspection"]
    end

    INPUTS --> COMP --> OUT
""", 6.5),

    # 9.6 Stage 0 Wire Format - Packet structure
    ("Packet (\u2264 1088 bytes total", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '15px'}}}%%
flowchart TB
    subgraph PKT["<b>Packet (&le; 1088 bytes total, big-endian)</b>"]
        direction TB
        subgraph HDR["<b>Header (32 B fixed)</b>"]
            direction TB
            H1["magic         4 B   0x52504144  ('RPAD')"]
            H2["version       1 B   0x01"]
            H3["packet_type   1 B   1=SESSION_START · 2=CHUNK · 3=SESSION_END"]
            H4["flags         2 B   bit 0: hmac_present"]
            H5["session_id    8 B   random 64-bit per submission"]
            H6["seq           4 B   monotonic uint32 within session"]
            H7["total_chunks  4 B   set in SESSION_START; 0 elsewhere"]
            H8["payload_len   2 B   length of bytes following"]
            H9["reserved      6 B   zeroed"]
            H1 --- H2 --- H3 --- H4 --- H5 --- H6 --- H7 --- H8 --- H9
        end
        PAY["<b>Payload (&le; 1024 B)</b>"]
        TAG["<b>HMAC tag (32 B)</b><br/>HMAC-SHA256 over (header || payload)<br/>using the per-drone shared secret"]
        HDR --> PAY --> TAG
    end

    style PKT fill:#FFF8E1,stroke:#F57F17,stroke-width:3px
    style HDR fill:#E3F2FD,stroke:#1565C0,stroke-width:2px
    style PAY fill:#E8F5E9,stroke:#388E3C,stroke-width:2px
    style TAG fill:#FFCDD2,stroke:#C62828,stroke-width:2px
""", 6.5),

    # 9.6 Application payload TLV substream
    ("Submission envelope (10 B)", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '15px'}}}%%
flowchart LR
    subgraph SUB["<b>Submission Envelope (10 B)</b>"]
        direction TB
        E["magic 4B 'DSUB'<br/>version 1B 0x01<br/>reserved 3B<br/>field_count 2B uint16"]
    end

    subgraph FLD["<b>TLV Record</b>"]
        direction TB
        F["tag 1B | length 2B | value NB<br/><i>(repeated field_count times)</i>"]
    end

    subgraph TLT["<b>Top-level Tags</b>"]
        direction TB
        TL["0x01 drone_id &mdash; utf-8<br/>0x02 timestamp &mdash; utf-8<br/>0x03 mission_id &mdash; utf-8 (opt)<br/>0x04 mission_zone &mdash; utf-8 (opt)<br/>0x05 geo &mdash; 24B (3 x float64)<br/>0x06 telemetry &mdash; 40B fixed<br/>0x07 signature &mdash; utf-8 (opt)<br/>0x08 firmware_version &mdash; utf-8 (opt)<br/>0x09 operator_id &mdash; utf-8 (opt)<br/>0x0A additional_metadata &mdash; TLV kv<br/>0x0B payload &mdash; TLV substream"]
    end

    subgraph PLT["<b>Payload Substream Tags</b>"]
        direction TB
        PL["0x01 type &mdash; 1B enum<br/>0x02 filename &mdash; utf-8<br/>0x03 mime &mdash; utf-8<br/>0x04 size_bytes &mdash; 8B uint64<br/>0x05 encryption &mdash; 1B bool<br/>0x06 container &mdash; 1B bool<br/>0x07 checksum &mdash; utf-8 (opt)<br/>0x08 uri &mdash; utf-8 (opt)"]
    end

    SUB --> FLD --> TLT
    TLT -.->|"0x0B references"| PLT

    style SUB fill:#E3F2FD,stroke:#1565C0,stroke-width:2px
    style FLD fill:#FFF9C4,stroke:#F9A825,stroke-width:2px
    style TLT fill:#E8F5E9,stroke:#388E3C,stroke-width:2px
    style PLT fill:#F3E5F5,stroke:#7B1FA2,stroke-width:2px
""", 6.5),

    # 16.1 Deployment Topology
    ("EDGE DEPLOYMENT NODE", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '18px'}}}%%
flowchart TB
    subgraph EDGE["EDGE DEPLOYMENT NODE"]
        subgraph IIP["Ingestion Interceptor Process"]
            PIPE["Interceptor<br/>Pipeline"]
            REG["Device<br/>Registry"]
            ULR["Uplink Receiver<br/>(gRPC/MQTT)"]
            LS[("Local Store<br/>(artifacts)")]
            PIPE --> LS
        end
        TE["Threat Estimator +<br/>Detection Engine"]
        IIP -->|"IngestResult"| TE
    end

    CC["Control Center /<br/>Security Dashboard"]
    ULR <-.->|"Secure Network Link"| CC

    style EDGE fill:#F3E5F5,stroke:#7B1FA2,stroke-width:2px
""", 5.5),
]


# ===================================================================
# THREAT ESTIMATOR DIAGRAM MAP
# ===================================================================

TE_DIAGRAM_MAP = [
    # 3.1 Pipeline position
    ("RPA / Drone Data Source", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}}}%%
flowchart LR
    DRONE["<b>Drone / RPA</b><br/>payload + metadata"]:::ext
    subgraph EDGE["EDGE MALWARE DETECTION ENGINE"]
        direction LR
        II["<b>Ingestion Interceptor</b><br/>auth + validate + extract"]
        TE["<b>Threat Estimator</b><br/>Stackelberg + Bayes<br/>emits T_S and<br/>inspection.level"]:::thismod
        subgraph DET["Multi-Layer Malware Detection"]
            direction TB
            SIG["Signature AV"]
            ML["AI/ML Classifier"]
            SBX["Sandbox Analyzer"]
        end
        MS["Metadata Sanitizer"]
        RM["Response &amp;<br/>Quarantine Manager"]
        DASH["Security Dashboard<br/>/ Feedback Loop"]
    end
    DRONE -->|"payload"| II
    II -->|"IngestResult"| TE
    TE -->|"T_S, inspection.route"| DET
    DET --> MS --> RM --> DASH
    RM -.->|"DetectionOutcome<br/>record_outcome()"| TE
    DASH -.->|"FPR / FNR<br/>update_thresholds_from_feedback()"| TE

    classDef thismod fill:#BBDEFB,stroke:#0D47A1,stroke-width:3px
    classDef ext fill:#E8F5E9,stroke:#388E3C,stroke-width:2px
    style EDGE fill:#FFF8E1,stroke:#F57F17,stroke-width:2px
    style DET fill:#E3F2FD,stroke:#1565C0,stroke-width:2px
""", 6.5),

    # 4.2 Component dependency (unique phrase: "(estimator.py)" appears only in this box)
    ("(estimator.py)", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}}}%%
flowchart TB
    TE["<b>GameTheoreticThreatEstimator</b><br/>(estimator.py)"]:::main

    BAY["<b>BayesianReputationEstimator</b><br/>(bayesian.py)<br/><i>cold-start prior P(N|Z,F)</i>"]
    SK["<b>Stackelberg helpers</b><br/>(stackelberg.py)<br/><i>I', DSR', payoffs, solver,<br/>T_S mapping</i>"]
    ADA["<b>AdaptiveThresholdManager</b><br/>(adaptive.py)<br/><i>soft FPR/FNR update</i>"]
    RS["<b>InMemoryReputationStore</b><br/>(reputation_store.py)<br/><i>swap for Redis/Postgres</i>"]
    FB["<b>FeedbackLoop</b><br/>(feedback.py)<br/><i>rolling FPR/FNR window</i>"]

    CFG["EstimatorConfig<br/>(config.py)"]:::cfg
    MOD["models.py<br/><i>ThreatEstimate, BayesianTrace,<br/>EquilibriumResult, …</i>"]:::mod

    TE --> BAY
    TE --> SK
    TE --> ADA
    TE --> RS
    TE --> FB

    CFG -.->|"tunables"| TE
    CFG -.->|"tunables"| ADA
    CFG -.->|"prior_benign"| BAY
    MOD -.->|"typed payloads"| TE

    classDef main fill:#BBDEFB,stroke:#0D47A1,stroke-width:3px
    classDef cfg fill:#FFF9C4,stroke:#F9A825,stroke-width:2px
    classDef mod fill:#E8F5E9,stroke:#388E3C,stroke-width:2px
""", 6.5),

    # 9.1 Happy path sequence (cold start)
    ("RepStore    Bayesian    Stackelberg   Adaptive", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '15px'}, 'sequence': {'mirrorActors': false}}}%%
sequenceDiagram
    autonumber
    participant C as Caller
    participant E as Estimator
    participant RS as RepStore
    participant B as Bayesian
    participant S as Stackelberg
    participant A as AdaptiveThresholds

    C->>E: estimate(ingest_metadata, artifacts)
    E->>RS: get(drone_id)
    RS-->>E: None (cold start)
    E->>B: compute_initial_reputation(Z, F)
    B-->>E: BayesianTrace (R)
    E->>E: compute I_base + flag bumps
    E->>S: compute_I_prime(I_base, R, Z)
    E->>S: compute_DSR_primes(DSR_base, H, TI)
    E->>S: build_payoff_matrices
    E->>S: solve_stackelberg_pure
    S-->>E: EquilibriumResult
    E->>S: compute_threat_score
    S-->>E: T_S
    E->>A: classify(T_S)
    A-->>E: InspectionDecision (level + route)
    E-->>C: ThreatEstimate
""", 6.5),

    # 9.2 Feedback cycle sequence
    ("record_outcome()", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '15px'}, 'sequence': {'mirrorActors': false}}}%%
sequenceDiagram
    autonumber
    participant H as Host
    participant E as Estimator
    participant RS as RepStore
    participant FB as FeedbackLoop
    participant A as AdaptiveThresholds

    rect rgb(240,248,255)
    Note over H,A: Detection outcome feedback
    H->>E: record_outcome(outcome)
    E->>RS: get(drone_id)
    RS-->>E: existing / None
    E->>E: update_reputation(current, verdict)
    E->>RS: put(new ReputationProfile)
    E->>FB: observe(outcome, inspection_level)
    E-->>H: ReputationProfile
    end

    rect rgb(255,248,240)
    Note over H,A: Periodic threshold refresh
    H->>E: update_thresholds_from_feedback()
    E->>FB: compute_metrics()
    FB-->>E: FeedbackMetrics (FPR, FNR)
    E->>A: update(fpr, fnr)
    A-->>E: (th_low, th_high)
    E-->>H: (th_low, th_high)
    end
""", 6.5),

    # 9.3 Rejection
    ("errored IngestResult", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}, 'sequence': {'mirrorActors': false}}}%%
sequenceDiagram
    autonumber
    participant C as Caller
    participant E as Estimator

    C->>E: estimate_from_ingest_result({"error": True, ...})
    E->>E: check shape — error flag set
    E-->>C: raise ValueError("cannot estimate on an errored IngestResult")
""", 6.0),
]


# ===================================================================
# SANDBOX ANALYZER DIAGRAM MAP
# ===================================================================

SB_DIAGRAM_MAP = [
    # 3.1 Pipeline position
    ("EDGE MALWARE DETECTION ENGINE", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}}}%%
flowchart TB
    DRONE["<b>Drone / RPA</b><br/>payload + metadata"]:::ext

    subgraph EDGE["EDGE MALWARE DETECTION ENGINE"]
        direction TB
        II["Ingestion Interceptor"]
        GTE["Game-Theoretic Threat Estimator"]
        ISS{"Inspection Strategy<br/>Selector"}:::decision

        subgraph LAYER3["Multi-Layer Malware Detection"]
            direction LR
            SIG["Signature Scanner<br/>(LOW)"]
            MLC["AI/ML Classifier<br/>(MEDIUM)"]
            SBX["<b>Sandbox Analyzer</b><br/>(HIGH)"]:::thismod
        end

        MS["Metadata Sanitizer"]
        TIC["Threat Intelligence<br/>Correlator"]
        RM["Response &amp;<br/>Quarantine Manager"]
        DASH["Security Dashboard<br/>&amp; Feedback Loop"]

        II --> GTE --> ISS
        ISS -->|"T_S LOW"| SIG
        ISS -->|"T_S MED"| MLC
        ISS -->|"T_S HIGH"| SBX
        LAYER3 --> MS --> TIC --> RM --> DASH
        RM -.->|"verdict feedback"| SBX
    end

    DRONE --> II

    classDef thismod fill:#BBDEFB,stroke:#0D47A1,stroke-width:3px
    classDef ext fill:#E8F5E9,stroke:#388E3C,stroke-width:2px
    classDef decision fill:#FFE0B2,stroke:#E65100,stroke-width:2px
    style EDGE fill:#FFF8E1,stroke:#F57F17,stroke-width:2px
    style LAYER3 fill:#E3F2FD,stroke:#1565C0,stroke-width:2px
""", 6.5),

    # 4.2 Internal 8-stage pipeline
    ("analyze(file_path, drone_id, ...)", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '15px'}}}%%
flowchart TB
    IN(["analyze(file_path, drone_id, …)"]):::in

    S1["<b>Stage 1 — File Router</b><br/>magic bytes + extension fallback"]:::stage
    SAFE(["SAFE → verdict CLEAN<br/>(short-circuit)"]):::clean
    WIN(["WINDOWS PE → verdict SUSPICIOUS<br/>(analyst review, never executed)"]):::susp

    S2["<b>Stage 2 — Archive Handler</b><br/>encrypted-no-key · double-ext · ZIP-bomb<br/>recursive extract (bounded)"]:::stage
    S3["<b>Stage 3 — Isolated Execution</b><br/>setrlimit preexec_fn · empty env<br/>watchdog (wall_timeout)"]:::stage
    S4["<b>Stage 4 — Four Monitors (200 ms poll)</b><br/>FileSystem · Network · Process · Privilege"]:::stage
    S5["<b>Stage 5 — Risk Scoring</b><br/>Σ behavior_weights + archive delta + IOC bonus"]:::stage
    S6["<b>Stage 6 — IOC Correlation</b><br/>ThreatIntelClient.query(hash, ips, domains)"]:::stage
    S7["<b>Stage 7 — Verdict</b><br/>CLEAN &lt; 25 · SUSPICIOUS &lt; 60 · MALICIOUS ≥ 60"]:::stage
    S8["<b>Stage 8 — FeedbackSink</b><br/>notify_response_manager · notify_dashboard · queue_for_ml"]:::stage

    OUT(["SandboxReport"]):::out

    IN --> S1
    S1 -->|"SAFE"| SAFE
    S1 -->|"WINDOWS"| WIN
    S1 -->|"ARCHIVE"| S2
    S1 -->|"SCRIPT / EXECUTABLE"| S3
    S2 --> S3
    S3 --> S4
    S4 --> S5 --> S6 --> S7 --> S8 --> OUT
    SAFE -.-> S8
    WIN -.-> S8

    classDef stage fill:#E3F2FD,stroke:#1565C0,stroke-width:2px
    classDef clean fill:#C8E6C9,stroke:#2E7D32,stroke-width:2px
    classDef susp fill:#FFE0B2,stroke:#E65100,stroke-width:2px
    classDef in fill:#FFF9C4,stroke:#F9A825,stroke-width:2px
    classDef out fill:#BBDEFB,stroke:#0D47A1,stroke-width:3px
""", 6.5),

    # 4.3 Component dependency
    ("(analyzer.py)", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}}}%%
flowchart TB
    SA["<b>SandboxAnalyzer</b><br/>(analyzer.py)<br/><i>orchestrates all 8 stages</i>"]:::main

    RT["route_file<br/>(router.py)"]
    AR["handle_archive<br/>(archive_handler.py)"]
    EX["execute_file<br/>(execution.py)"]
    MON["FileSystem / Network /<br/>Process / Privilege Monitor<br/>(monitors.py)"]
    SC["compute_risk_score<br/>(scoring.py)"]
    VER["determine_verdict<br/>(verdict.py)"]

    IOC["<b>ThreatIntelClient</b><br/>(ioc.py)<br/><i>LocalThreatIntelClient → swap<br/>for Correlator / VT / OTX / MISP</i>"]:::seam
    FB["<b>FeedbackSink</b><br/>(feedback.py)<br/><i>LoggingFeedbackSink → bridge to<br/>Response Mgr / Dashboard / ML</i>"]:::seam

    CFG["SandboxConfig<br/>(config.py)"]:::cfg
    MOD["models.py<br/><i>SandboxReport, MonitorEvent,<br/>ScoreBreakdown, …</i>"]:::mod

    SA --> RT
    SA --> AR
    SA --> EX
    SA --> MON
    SA --> SC
    SA --> VER
    SA --> IOC
    SA --> FB

    CFG -.->|"tunables"| SA
    MOD -.->|"typed payloads"| SA

    classDef main fill:#BBDEFB,stroke:#0D47A1,stroke-width:3px
    classDef seam fill:#FFE0B2,stroke:#E65100,stroke-width:2px
    classDef cfg fill:#FFF9C4,stroke:#F9A825,stroke-width:2px
    classDef mod fill:#E8F5E9,stroke:#388E3C,stroke-width:2px
""", 6.5),

    # 9.1 Happy path SAFE CSV
    ("SAFE, should_skip=True", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}, 'sequence': {'mirrorActors': false}}}%%
sequenceDiagram
    autonumber
    participant C as Caller
    participant A as SandboxAnalyzer
    participant R as Router
    participant F as FeedbackSink

    C->>A: analyze(file_path, drone_id)
    A->>R: route_file()
    R-->>A: SAFE · should_skip=True
    A->>A: short-circuit — verdict = CLEAN
    A->>F: notify_response_manager / notify_dashboard / queue_for_ml
    A-->>C: SandboxReport
""", 6.5),

    # 9.2 Archive path
    ("ArchiveHandler    Scoring   Verdict   Feedback", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '14px'}, 'sequence': {'mirrorActors': false}}}%%
sequenceDiagram
    autonumber
    participant C as Caller
    participant A as SandboxAnalyzer
    participant R as Router
    participant H as ArchiveHandler
    participant S as Scoring
    participant TI as TIClient
    participant V as Verdict
    participant F as FeedbackSink

    C->>A: analyze(zip, drone_id)
    A->>R: route_file()
    R-->>A: ARCHIVE
    A->>H: handle_archive(zip, out, config)
    H->>H: scan listing — double-ext detected
    H-->>A: ArchiveResult (+30 risk)
    A->>A: extract + per-file execute (Stages 3-4)
    A->>S: compute_risk_score(events)
    S-->>A: ScoreBreakdown
    A->>TI: ThreatIntelClient.query(hash, ips, domains)
    TI-->>A: IOCResult
    A->>V: determine_verdict(total_score)
    V-->>A: SUSPICIOUS / MALICIOUS
    A->>F: dispatch feedback (3 channels)
    A-->>C: SandboxReport
""", 6.5),

    # 9.3 Execute + monitor
    ("TI-Client   Scoring   FeedbackSink", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '14px'}, 'sequence': {'mirrorActors': false}}}%%
sequenceDiagram
    autonumber
    participant C as Caller
    participant A as SandboxAnalyzer
    participant R as Router
    participant E as execute_file
    participant M as Monitors
    participant TI as TIClient
    participant S as Scoring
    participant F as FeedbackSink

    C->>A: analyze(script, drone_id)
    A->>R: route_file()
    R-->>A: SCRIPT / EXECUTABLE
    A->>E: subprocess.Popen + preexec setrlimit
    E->>A: on_pid(pid)
    A->>M: build_monitors(pid, log, config)
    loop every poll_interval_seconds (default 200 ms)
        M->>M: poll /proc + psutil
    end
    E-->>A: stdout / stderr / exit_code
    A->>S: compute_risk_score(events)
    A->>TI: query(file_hash, ips, domains)
    TI-->>A: IOCResult
    A->>F: dispatch feedback
    A-->>C: SandboxReport
""", 6.5),
]

MS_DIAGRAM_MAP = [
    # 3.1 Position in Detection Pipeline
    ("INGESTION INTERCEPTOR", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}}}%%
flowchart TB
    DRONE["<b>DRONE / RPA</b>"]

    subgraph EDGE["EDGE DEPLOYMENT NODE<br/><i>(all modules run in-process)</i>"]
        direction TB
        II["<b>Ingestion Interceptor</b>"]
        GTE["<b>Game-Theoretic Estimator</b><br/>T_S (0-1)"]
        ISS["<b>Inspection Strategy<br/>Selector</b>"]
        MDE["<b>Malware Detection</b><br/>Sig + ML + Sandbox"]
        subgraph MS["<b>METADATA SANITIZER</b>"]
            direction LR
            IH["Image"] --- VH["Video"] --- PH["PDF"]
            AH["Archive"] --- TH["Text"]
        end
        TIC["<b>Threat Intel Correlator</b>"]
        RQM["<b>Response &amp; Quarantine</b>"]
        DASH["<b>Security Dashboard</b>"]

        II -->|"IngestResult"| GTE -->|"T_S"| ISS -->|"inspection level"| MDE -->|"passed"| MS -->|"cleaned"| TIC --> RQM --> DASH
        DASH -.->|"Uplink cmds"| II
    end

    GCC["<b>Ground Control<br/>Center</b>"]:::gcc

    DRONE -->|"native protocol /<br/>Stage 0 UDP"| II
    DASH <-->|"packets<br/>(MQTT / gRPC over TLS)"| GCC

    style EDGE fill:#FFF8E1,stroke:#F57F17,stroke-width:2px
    style MS fill:#FFF9C4,stroke:#F9A825,stroke-width:2px
    classDef gcc fill:#E8F5E9,stroke:#388E3C,stroke-width:2px
""", 6.5),

    # 3.3 Data Flow Between Modules
    ("ArtifactRecord[]", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}}}%%
flowchart LR
    II["Ingestion<br/>Interceptor"] -->|"IngestMetadata +<br/>ArtifactRecord[]"| GTE["Game-Theoretic<br/>Estimator"]
    GTE -->|"T_S"| MDE["Malware<br/>Detection"]
    MDE -->|"passed files"| MSB

    subgraph MSB["Metadata Sanitizer"]
        direction TB
        INP["<b>Inputs:</b> artifact_record, T_S, insecure_flags"]
        OUTP["<b>Outputs:</b> SanitizationResult + cleaned file"]
    end

    MSB -->|"report"| TIC["Threat Intel<br/>Correlator"]
    MSB -->|"cleaned file"| ON["Operational<br/>Network"]
""", 6.5),

    # 4.2 Component Diagram (MetadataSanitizer)
    ("MetadataSanitizer", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}}}%%
flowchart TB
    subgraph MS["<b>MetadataSanitizer</b> (sanitizer.py)"]
        direction TB
        MR["Mode<br/>Resolver"] & SPR["Storage Pointer<br/>Resolver"] & ST["Statistics<br/>Tracker"]
        subgraph HR["Handler Registry"]
            direction LR
            MIME["MIME type"] --> HC["Handler class"] --> HI["cached instance"]
        end
        MR & SPR --> HR
        subgraph HL["Handlers"]
            direction LR
            IH["Image"] --- VH["Video"] --- PH["PDF"] --- AH["Archive"] --- TH["Text"]
        end
        HR --> HL
        RULES["Sanitization Rules<br/>exif_rules | pdf_rules | video_rules"]
        HL --> RULES
    end
    SRES["SanitizationResult<br/>(audit trail)"]
    CFILE["Cleaned File (on disk)"]
    MS --> SRES & CFILE
""", 6.0),

    # 4.3 Handler Interface (class diagram)
    ("BaseHandler (ABC)", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}}}%%
classDiagram
    class BaseHandler {
        <<abstract>>
        +config: SanitizerConfig
        +logger: Logger
        +extract_metadata(file_path) Dict
        +sanitize(file, output, mode) Changes[]
        +verify(file_path) bool
        +supported_mimes() Set
        +is_available() bool
        +handler_name() str
        #create_metadata_snapshot(meta) Snapshot
        #make_change(field, action) Change
        #check_field_size_anomaly(field, val, th)
    }
    class ImageHandler { EXIF scrubbing via Pillow+piexif }
    class VideoHandler { Atom cleaning via mutagen }
    class PdfHandler { JS/action removal via pikepdf }
    class ArchiveHandler { Structure inspection (stdlib) }
    class TextHandler { Encoding normalization (stdlib) }

    BaseHandler <|-- ImageHandler
    BaseHandler <|-- VideoHandler
    BaseHandler <|-- PdfHandler
    BaseHandler <|-- ArchiveHandler
    BaseHandler <|-- TextHandler
""", 6.0),

    # 5.1 Sanitization Pipeline
    ("1. Resolve Mode", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}}}%%
flowchart TB
    IN(["Input: artifact_record + threat_score"])

    subgraph ROW1[" "]
        direction LR
        S1["<b>1. Resolve Mode</b><br/>threat_score &rarr; mode"]
        S2["<b>2. Pre-Checks</b><br/>file? size? MIME?"]
        S3["<b>3. Extract Before</b><br/>MetadataSnapshot"]
        S4["<b>4. Preserve</b><br/>copy &rarr; .orig"]
        S1 --> S2 --> S3 --> S4
    end

    subgraph ROW2[" "]
        direction LR
        S5["<b>5. Sanitize</b><br/>handler.sanitize()"]
        S6["<b>6. Verify</b><br/>re-parse output"]
        S7["<b>7. Extract After</b><br/>compare snapshots"]
        S5 --> S6 --> S7
    end

    OUT(["SanitizationResult"])
    IN --> ROW1 --> ROW2 --> OUT

    style ROW1 fill:none,stroke:none
    style ROW2 fill:none,stroke:none
""", 6.0),

    # 5.2 Mode Resolver
    ("Priority 1: mode_override", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}}}%%
flowchart TB
    P1{"Priority 1: mode_override provided?"}
    P1 -->|"Yes"| USE1(["Use directly"])
    P1 -->|"No"| P2{"Priority 2: threat_score available?"}

    P2 -->|"T_S &ge; 0.7"| STRIP["<b>strip</b>"]:::strip
    P2 -->|"T_S &le; 0.3"| AUDIT["<b>audit_only</b>"]:::audit
    P2 -->|"0.3 &lt; T_S &lt; 0.7"| SEL["<b>selective</b>"]:::sel
    P2 -->|"No T_S"| P3{"Priority 3: insecure_flags?"}

    P3 -->|"executable / suspicious /<br/>double_extension"| STRIP2["<b>strip</b>"]:::strip
    P3 -->|"No critical flags"| P4["Priority 4: config.default_mode<br/>&rarr; <b>selective</b>"]:::sel

    classDef strip fill:#FFCDD2,stroke:#C62828
    classDef audit fill:#C8E6C9,stroke:#2E7D32
    classDef sel fill:#FFF9C4,stroke:#F9A825
""", 5.5),

    # 5.4 Image Handler Flow
    ("piexif available", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}}}%%
flowchart LR
    IMG(["Image File"]) --> CHECK{"piexif?"}

    CHECK -->|"YES"| YES
    CHECK -->|"NO"| NO

    subgraph YES["Surgical Removal"]
        direction LR
        Y1["Load EXIF"] --> Y2["Match rules"] --> Y3["Delete tags"] --> Y4["Re-insert"]
    end

    subgraph NO["Pillow Fallback"]
        direction LR
        N1["Copy pixels"] --> N2["Re-save<br/>without EXIF"]
    end

    style YES fill:#E8F5E9,stroke:#388E3C
    style NO fill:#FFF3E0,stroke:#EF6C00
""", 6.5),

    # 5.5 PDF Handler
    ("Strip catalog keys", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}}}%%
flowchart TB
    PDF(["PDF File"])
    subgraph STEPS[" "]
        direction LR
        S1["<b>1. Strip Catalog</b><br/>/JavaScript, /OpenAction,<br/>/Launch, /SubmitForm,<br/>/EmbeddedFiles, /XFA"]
        S2["<b>2. Strip Page Actions</b><br/>/AA per page,<br/>annotation actions"]
        S3["<b>3. Clean DocInfo</b><br/>strip: clear all<br/>selective: safe keys"]
        S4["<b>4. Scan Streams</b><br/>JS, shellcode, heap spray,<br/>PowerShell, PE headers"]
        S1 --> S2 --> S3 --> S4
    end
    PDF --> STEPS
    style STEPS fill:none,stroke:none
""", 6.0),

    # 8.1 Single File Sanitization Flow (sequence)
    ("Client              MetadataSanitizer        Handler", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}, 'sequence': {'mirrorActors': false}}}%%
sequenceDiagram
    participant C as Client
    participant MS as MetadataSanitizer
    participant H as Handler
    participant R as Rules
    participant FS as FileSystem

    C->>MS: sanitize_file(...)
    MS->>MS: resolve_mode() (threat_score -> mode)
    MS->>MS: pre-checks (file? size? MIME?)
    MS->>H: get_handler()
    H-->>MS: handler instance
    MS->>H: extract_metadata()
    H->>FS: read file
    FS-->>H: file bytes
    H-->>MS: before_snapshot
    MS->>FS: copy -> .orig (preserve original)
    MS->>H: sanitize()
    H->>R: get_strip_set()
    R-->>H: tag set
    H->>FS: write cleaned file
    H-->>MS: changes[]
    MS->>H: verify()
    H->>FS: re-parse file
    H-->>MS: valid=true
    MS-->>C: SanitizationResult
""", 6.5),

    # 8.2 Threat-Score-Driven Mode Selection
    ("T_S (Threat Score)", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '18px'}}}%%
flowchart TB
    TS(["T_S (Threat Score)"])

    TS -->|"T_S &le; 0.3"| AO
    TS -->|"0.3 &lt; T_S &lt; 0.7"| SEL
    TS -->|"T_S &ge; 0.7"| STR

    subgraph AO["AUDIT_ONLY"]
        AO_D["Log only, no modifications<br/><br/><i>Low threat (routine patrol)</i>"]
    end

    subgraph SEL["SELECTIVE"]
        SEL_D["Remove known-dangerous<br/>fields only<br/><br/><i>Medium threat (default)</i>"]
    end

    subgraph STR["STRIP"]
        STR_D["Remove ALL non-essential<br/>metadata<br/><br/><i>High threat (suspicious device)</i>"]
    end

    style AO fill:#C8E6C9,stroke:#2E7D32
    style SEL fill:#FFF9C4,stroke:#F9A825
    style STR fill:#FFCDD2,stroke:#C62828
""", 5.5),

    # 11.2 Defense-in-Depth Layers
    ("Layer 1: Ingestion Interceptor", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '16px'}}}%%
flowchart TB
    subgraph ROW1[" "]
        direction LR
        L1["<b>Layer 1:</b> Ingestion Interceptor<br/><i>JSON validation, auth</i>"]
        L2["<b>Layer 2:</b> Game-Theoretic Estimator<br/><i>Threat score</i>"]
        L3["<b>Layer 3:</b> Malware Detection<br/><i>Sig + ML + Sandbox</i>"]
        L1 --> L2 --> L3
    end
    subgraph ROW2[" "]
        direction LR
        L4["<b>Layer 4: METADATA SANITIZER</b><br/><i>Metadata cleaning, scripts</i>"]:::hl
        L5["<b>Layer 5:</b> Threat Intel Correlator<br/><i>IoC matching</i>"]
        L6["<b>Layer 6:</b> Response Manager<br/><i>Quarantine / forward</i>"]
        L4 --> L5 --> L6
    end
    ROW1 --> ROW2

    classDef hl fill:#FFF9C4,stroke:#F9A825,stroke-width:3px
    style ROW1 fill:none,stroke:none
    style ROW2 fill:none,stroke:none
""", 6.0),

    # 15.1 Deployment Topology
    ("Edge Node", r"""
%%{init: {'theme': 'base', 'themeVariables': {'fontSize': '18px'}}}%%
flowchart LR
    subgraph EDGE["Edge Node"]
        direction LR
        II["Ingestion Interceptor"]
        MS["Metadata Sanitizer"]
        SHARED[("drone_remote_store/")]
        II --> MS --> SHARED
    end
""", 5.5),
]


# ===================================================================
# RENDERING
# ===================================================================

def render_mermaid(name: str, code: str) -> str:
    """Render Mermaid code to PNG, return path or empty string."""
    mmd_path = os.path.join(DIAGRAMS_DIR, f"{name}.mmd")
    png_path = os.path.join(DIAGRAMS_DIR, f"{name}.png")

    with open(mmd_path, "w") as f:
        f.write(code.strip())

    pup_cfg = os.path.join(DIAGRAMS_DIR, "puppeteer-config.json")
    if not os.path.exists(pup_cfg):
        with open(pup_cfg, "w") as f:
            json.dump({"args": ["--no-sandbox", "--disable-setuid-sandbox"]}, f)

    cmd = ["mmdc", "-i", mmd_path, "-o", png_path,
           "-w", "3000", "-s", "2.5", "-b", "white", "-p", pup_cfg]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
    if r.returncode != 0:
        print(f"  WARN: {name}: {r.stderr[:200]}")
        return ""
    print(f"  OK: {name}.png")
    return png_path


def render_all_diagrams():
    """Render all Mermaid diagrams for all design documents."""
    rendered = {}
    print("Rendering Ingestion Interceptor diagrams...")
    for i, (match, code, width) in enumerate(II_DIAGRAM_MAP):
        name = f"ii_{i:02d}"
        path = render_mermaid(name, code)
        if path:
            rendered[("ii", match)] = (path, width)

    print("Rendering Metadata Sanitizer diagrams...")
    for i, (match, code, width) in enumerate(MS_DIAGRAM_MAP):
        name = f"ms_{i:02d}"
        path = render_mermaid(name, code)
        if path:
            rendered[("ms", match)] = (path, width)

    print("Rendering Threat Estimator diagrams...")
    for i, (match, code, width) in enumerate(TE_DIAGRAM_MAP):
        name = f"te_{i:02d}"
        path = render_mermaid(name, code)
        if path:
            rendered[("te", match)] = (path, width)

    print("Rendering Sandbox Analyzer diagrams...")
    for i, (match, code, width) in enumerate(SB_DIAGRAM_MAP):
        name = f"sb_{i:02d}"
        path = render_mermaid(name, code)
        if path:
            rendered[("sb", match)] = (path, width)

    return rendered


# ===================================================================
# MARKDOWN PARSER + DOCX BUILDER
# ===================================================================

def setup_styles(doc):
    """Configure document styles for professional appearance."""
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.space_before = Pt(2)

    # Configure heading styles
    for level in range(1, 5):
        hstyle = doc.styles[f"Heading {level}"]
        hstyle.font.name = "Calibri"
        hstyle.font.color.rgb = RGBColor(0x1A, 0x23, 0x7E)
        if level == 1:
            hstyle.font.size = Pt(22)
            hstyle.paragraph_format.space_before = Pt(24)
            hstyle.paragraph_format.space_after = Pt(12)
        elif level == 2:
            hstyle.font.size = Pt(17)
            hstyle.paragraph_format.space_before = Pt(20)
            hstyle.paragraph_format.space_after = Pt(8)
        elif level == 3:
            hstyle.font.size = Pt(14)
            hstyle.paragraph_format.space_before = Pt(14)
            hstyle.paragraph_format.space_after = Pt(6)
        elif level == 4:
            hstyle.font.size = Pt(12)
            hstyle.paragraph_format.space_before = Pt(10)
            hstyle.paragraph_format.space_after = Pt(4)

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)


def add_formatted_run(para, text):
    """Add text to paragraph, handling **bold**, *italic*, `code` inline."""
    # Split into segments of formatted and plain text
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`[^`]+`)', text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = para.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and not part.startswith("**"):
            run = para.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = para.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        else:
            if part:
                para.add_run(part)


def add_code_block(doc, code_text, language=""):
    """Add a code block with grey background and monospace font."""
    # Add a thin border table as a container for the code block
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    cell = table.cell(0, 0)
    # Set cell shading to light grey
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F5F5F5"/>')
    cell._tc.get_or_add_tcPr().append(shading)

    # Set cell border
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '  <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '  <w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '  <w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        '</w:tcBorders>'
    )
    tc_pr.append(borders)

    # Set cell width
    cell.width = Inches(6.5)

    # Clear existing content and add code
    cell.text = ""
    lines = code_text.rstrip("\n").split("\n")
    for i, line in enumerate(lines):
        if i == 0:
            p = cell.paragraphs[0]
        else:
            p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(14)
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    doc.add_paragraph()  # spacing after


def add_md_table(doc, header_row, data_rows):
    """Add a formatted table to the document."""
    ncols = len(header_row)
    table = doc.add_table(rows=1 + len(data_rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # Header
    hdr = table.rows[0]
    for i, cell_text in enumerate(header_row):
        cell = hdr.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(cell_text.strip())
        run.bold = True
        run.font.size = Pt(10)
        run.font.name = "Calibri"
        # Blue header background
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E3F2FD"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    # Data rows
    for ri, row_data in enumerate(data_rows):
        row = table.rows[ri + 1]
        for ci in range(ncols):
            cell_text = row_data[ci].strip() if ci < len(row_data) else ""
            cell = row.cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            add_formatted_run(p, cell_text)
            for run in p.runs:
                run.font.size = Pt(10)
                run.font.name = "Calibri"
            # Alternate row shading
            if ri % 2 == 1:
                shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FAFAFA"/>')
                cell._tc.get_or_add_tcPr().append(shading)

    doc.add_paragraph()  # spacing


def parse_table_lines(lines):
    """Parse markdown table lines into header and data rows."""
    rows = []
    for line in lines:
        line = line.strip()
        if line.startswith("|") and line.endswith("|"):
            cells = [c.strip() for c in line.split("|")[1:-1]]
            # Skip separator rows (---|---|)
            if cells and all(re.match(r'^[-:]+$', c) for c in cells):
                continue
            rows.append(cells)
    if len(rows) >= 2:
        return rows[0], rows[1:]
    elif len(rows) == 1:
        return rows[0], []
    return None, None


def is_ascii_art(text):
    """Check if text contains box-drawing characters or diagram-like arrows."""
    has_box = any(c in text for c in "┌┐└┘├┤┬┴─│▼▲►◄△╔╗╚╝═║")
    # Also detect text-based diagrams with arrows and indentation
    has_arrow_diagram = ("-->" in text or "──►" in text or "──>" in text) and \
                        text.count("\n") > 3 and \
                        re.search(r'^\s{2,}', text, re.MULTILINE) is not None
    return has_box or has_arrow_diagram


def find_matching_diagram(block_text, diagram_map, prefix):
    """Find the diagram that matches this block's content."""
    for match_str, code, width in diagram_map:
        if match_str in block_text:
            return (match_str, code, width)
    return None


def add_image_to_doc(doc, png_path, width_inches, caption=""):
    """Add a centered image with optional caption, capping height to fit page."""
    from PIL import Image as PILImage
    max_height_inches = 8.0  # keep within a single page

    img = PILImage.open(png_path)
    img_w, img_h = img.size
    aspect = img_h / img_w

    target_w = width_inches
    target_h = target_w * aspect

    # If too tall, scale down to fit max height
    if target_h > max_height_inches:
        target_h = max_height_inches
        target_w = target_h / aspect

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(png_path, width=Inches(target_w), height=Inches(target_h))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        run.font.size = Pt(9)
        run.font.italic = True
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def add_blockquote(doc, bq_lines, add_formatted_run_fn, add_md_table_fn, parse_table_lines_fn):
    """Render a blockquote block as an indented, left-bordered callout."""
    # Strip '> ' prefix from each line
    stripped = []
    for l in bq_lines:
        if l.startswith('> '):
            stripped.append(l[2:])
        elif l.strip() == '>':
            stripped.append('')
        else:
            stripped.append(l)

    # Create a single-cell table as the callout container
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.3)

    # Light blue background + blue left border
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EBF5FB"/>')
    tc_pr.append(shading)
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '  <w:top w:val="single" w:sz="4" w:space="0" w:color="D6EAF8"/>'
        '  <w:left w:val="single" w:sz="16" w:space="0" w:color="2E86C1"/>'
        '  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D6EAF8"/>'
        '  <w:right w:val="single" w:sz="4" w:space="0" w:color="D6EAF8"/>'
        '</w:tcBorders>'
    )
    tc_pr.append(borders)

    # Clear the default empty paragraph
    cell.text = ""

    # Parse content inside the blockquote
    j = 0
    first_para = True
    while j < len(stripped):
        sl = stripped[j]

        # Blank line
        if sl.strip() == "":
            j += 1
            continue

        # Table inside blockquote
        if sl.strip().startswith("|") and "|" in sl.strip()[1:]:
            tbl_lines = []
            while j < len(stripped) and stripped[j].strip().startswith("|"):
                tbl_lines.append(stripped[j])
                j += 1
            header, rows = parse_table_lines_fn(tbl_lines)
            if header and rows:
                # Render the table OUTSIDE the callout cell (nested tables are messy)
                # Instead, render as formatted text rows inside the cell
                # Header row
                p = cell.paragraphs[0] if first_para else cell.add_paragraph()
                first_para = False
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(" | ".join(header))
                run.bold = True
                run.font.size = Pt(9)
                run.font.name = "Calibri"
                # Data rows
                for row_data in rows:
                    p = cell.add_paragraph()
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    row_text = " | ".join(c.strip() for c in row_data)
                    add_formatted_run_fn(p, row_text)
                    for run in p.runs:
                        run.font.size = Pt(9)
                        run.font.name = "Calibri"
            continue

        # Unordered list inside blockquote
        if re.match(r'^[-*]\s+', sl):
            while j < len(stripped) and re.match(r'^[-*]\s+', stripped[j]):
                text = re.sub(r'^[-*]\s+', '', stripped[j])
                p = cell.paragraphs[0] if first_para else cell.add_paragraph()
                first_para = False
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Inches(0.2)
                bullet_run = p.add_run("\u2022 ")
                bullet_run.font.size = Pt(10)
                add_formatted_run_fn(p, text)
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.font.name = "Calibri"
                j += 1
            continue

        # Regular paragraph inside blockquote
        para_parts = []
        while j < len(stripped) and stripped[j].strip() != "" \
              and not stripped[j].strip().startswith("|") \
              and not re.match(r'^[-*]\s+', stripped[j]):
            para_parts.append(stripped[j])
            j += 1
        if para_parts:
            text = " ".join(para_parts)
            p = cell.paragraphs[0] if first_para else cell.add_paragraph()
            first_para = False
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            add_formatted_run_fn(p, text)
            for run in p.runs:
                run.font.size = Pt(10)
                run.font.name = "Calibri"
                run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    doc.add_paragraph()  # spacing after callout


def parse_box_table(text):
    """Try to parse ASCII art box table into header + rows."""
    lines = text.strip().split("\n")
    title = None
    header = None
    rows = []
    footer_lines = []
    in_footer = False

    for line in lines:
        stripped = line.strip()
        # Skip border-only lines
        if re.match(r'^[┌┐└┘├┤┬┴─┼╔╗╚╝═╬]+$', stripped):
            continue
        # Extract cell content
        if "│" in stripped:
            cells = [c.strip() for c in stripped.split("│")]
            cells = [c for c in cells if c]  # remove empty from edges
            if not cells:
                continue
            # Check if it's a single-cell line (title or footer)
            if len(cells) == 1:
                if header is None and not rows:
                    title = cells[0]
                else:
                    footer_lines.append(cells[0])
            else:
                if header is None:
                    header = cells
                else:
                    rows.append(cells)

    if header and rows:
        return title, header, rows, footer_lines
    return None, None, None, None


# ===================================================================
# SECTION INTRO DESCRIPTIONS
# ===================================================================

II_SECTION_DESCS = {
    "2. Goals and Non-Goals": "This section defines the precise scope of the Ingestion Interceptor by listing what the module is designed to achieve and what responsibilities are explicitly delegated to other modules in the pipeline.",
    "2.1 Goals": "The following goals define the core functional responsibilities of the Ingestion Interceptor.",
    "2.2 Non-Goals": "The following items are explicitly outside the scope of this module and are handled by downstream components in the detection pipeline.",
    "3.1 Full System Architecture": "The following diagram illustrates the complete three-layer system architecture, showing all nine modules and the data flow between them.",
    "3.2 Ingestion Interceptor Boundary Context": "The following diagram shows the trust boundary in which the Ingestion Interceptor operates, highlighting external inputs from the drone platform (UDP packets at Stage 0 or in-process dict via process()) and the security dashboard control center.",
    "4. Architecture Overview": "This section describes the internal architecture of the Ingestion Interceptor module, including its file structure, component dependencies, and the data flow through its eight-stage processing pipeline (Stage 0 packet receiver plus the original seven stages).",
    "4.1 Module Structure": "The module is organized as a Python package with each processing stage implemented in a separate file for clear separation of concerns. The following listing shows the directory layout, including the new Stage 0 packet receiver and the protocol/ wire-codec subpackage.",
    "4.2 Component Dependency Diagram": "The following diagram shows how the main orchestrator (interceptor.py) depends on each processing component, how the shared configuration object flows to all modules, and how the Stage 0 packet_receiver and protocol/ subpackage feed reassembled submissions back into the orchestrator.",
    "4.3 Data Flow Through the Pipeline": "The following diagram traces the complete data flow, beginning at the optional Stage 0 UDP packet receiver and continuing through the seven processing stages to the final IngestResult output, including the rejection paths.",
    "5. Component Design": "This section provides detailed design documentation for each component in the Ingestion Interceptor module, describing their purpose, internal architecture, validation logic, and key design decisions.",
    "5.10 `packet_receiver.py`": "This component implements Stage 0: a UDP listener thread, a per-session reassembly buffer, periodic garbage collection, and packet/session statistics. It is the wire-level entry point for drone submissions.",
    "5.11 `protocol/`": "The protocol subpackage holds the wire contract: the 32-byte fixed packet header, the per-packet HMAC tag, and the binary TLV submission marshaller. Both the receiver and the drone simulator import from this subpackage so the wire contract has a single source of truth.",
    "6. Data Models": "This section defines the typed dataclass models used throughout the pipeline, from input parsing to final output assembly. Each model includes field definitions with types, descriptions, and serialization methods.",
    "7. API Specification": "This section documents the public API surface of the Ingestion Interceptor, including both the class-based and functional entry points, along with module-level utility functions for each component.",
    "8. Sequence Diagrams": "This section provides sequence diagrams that illustrate the interaction between components during key processing scenarios, including the happy path, authentication flow, uplink command handling, Stage 0 packet reassembly, and rejection handling.",
    "8.4 Stage 0": "The following sequence diagrams illustrate the wire-level Stage 0 lifecycle: the happy path of fragmenting, transmitting, verifying, and reassembling a submission via UDP, followed by the rejection and truncation paths that surface as PacketReceiverStats counters.",
    "9. Input/Output Specification": "This section documents the input formats accepted by the Ingestion Interceptor (the in-process JSON dict and the wire-level binary protocol used by Stage 0) along with the output IngestResult structure for the success, error, and flagged paths.",
    "9.6 Stage 0 Wire Format": "The following diagrams describe the binary wire format used by Stage 0: the packet framing (32-byte header, payload, HMAC tag) and the TLV-encoded submission envelope carried inside CHUNK packets.",
    "10. Security Considerations": "This section documents the security defences implemented across all pipeline stages, including the new Stage 0 wire-layer defenses, mapping each defence mechanism to the specific threat it mitigates.",
    "11. Configuration Reference": "This section provides the complete reference for all tuneable parameters in the InterceptorConfig dataclass, organized by functional area with defaults and descriptions.",
}

TE_SECTION_DESCS = {
    "2. Goals and Non-Goals": "This section defines the scope of the Game-Theoretic Threat Estimator by listing the capabilities it is responsible for and the responsibilities explicitly delegated to other modules.",
    "2.1 Goals": "The following goals define the core functional responsibilities of the Threat Estimator within the edge malware detection pipeline.",
    "2.2 Non-Goals": "The following items are deliberately outside the scope of this module.",
    "3. System Context": "This section places the Threat Estimator in the broader edge malware detection pipeline and identifies the four dependency-injected seams through which it connects to production backends.",
    "3.1 Position in the edge malware detection pipeline": "The following diagram shows how the Threat Estimator sits between the Ingestion Interceptor and the Multi-Layer Malware Detection Engine, and how verdicts from the Response Manager and metrics from the Dashboard feed back into the estimator.",
    "3.2 Integration seams": "The following table enumerates the four dependency-injected integration points; each has a reference implementation that is safe in isolation and a clear path to a production backend.",
    "4. Architecture Overview": "This section describes the internal architecture of the Threat Estimator module: its file layout, component dependencies, and the data flow of a single estimate.",
    "4.1 Module layout": "The module is organised as a flat Python package with each algorithmic concern in its own file.",
    "4.2 Component dependency diagram": "The following diagram shows how GameTheoreticThreatEstimator composes the Bayesian estimator, Stackelberg helpers, adaptive threshold manager, reputation store, and feedback loop.",
    "4.3 Data flow through the pipeline": "The following numbered steps trace a single call to estimate() from raw IngestResult to an emitted ThreatEstimate.",
    "5. Component Design": "This section provides detailed design documentation for every component, including the orchestrator, algorithmic helpers, persistent stores, and feedback-loop aggregator.",
    "6. Data Models": "This section defines the typed dataclasses exchanged across the Threat Estimator's public API, including the ThreatEstimate output and its supporting models.",
    "7. API Specification": "This section documents the public Python API surface: constructor arguments, primary methods, and an end-to-end integration example.",
    "8. Input / Output Specification": "This section defines the exact input contract accepted by estimate() and the stable output contract emitted as ThreatEstimate.to_dict(), plus the downstream consumption contract per module.",
    "9. Sequence Diagrams": "This section provides sequence diagrams for the three canonical flows: a cold-start happy-path estimate, the verdict-feedback + threshold-refresh cycle, and the rejection of an errored IngestResult.",
    "10. Mathematical Core": "This section documents every formula used by the estimator: adjusted impact, adjusted detection success rates, the Stackelberg payoffs and solver, the sigmoid-blend threat-score mapping, the adaptive-threshold update rule, the Bayesian cold-start posterior, and the asymmetric reputation update.",
    "11. Configuration Reference": "This section enumerates every tunable parameter on EstimatorConfig with its default value and meaning, grouped by functional area.",
    "12. Security & Threat Model": "This section enumerates the adversary model, describes the non-repudiation guarantees, specifies the input trust boundary, and documents the fail-closed configuration validation.",
    "13. Performance Characteristics": "This section documents the latency budget per phase, the memory bounds of every internal structure, and the expected throughput per core.",
    "14. Testing Strategy": "This section describes the 48-test suite layout, how to run it, and the invariants explicitly enforced by the tests.",
    "15. Operations": "This section covers deployment, the operational monitoring checklist, and runbook snippets for common incidents.",
    "16. Risk Assessment": "This section enumerates known risks with their likelihood, impact, and the in-module mitigation.",
}

SB_SECTION_DESCS = {
    "2. Goals and Non-Goals": "This section defines the scope of the Sandbox Analyzer by listing the capabilities it implements and the responsibilities explicitly handled by other modules.",
    "2.1 Goals": "The following goals define the dynamic-analysis responsibilities of the sandbox.",
    "2.2 Non-Goals": "The following items are explicitly outside the scope of this module.",
    "3. System Context": "This section places the Sandbox Analyzer inside the Multi-Layer Malware Detection Engine and enumerates its integration seams.",
    "3.1 Position in the edge malware detection pipeline": "The following diagram shows how the Inspection Strategy Selector routes HIGH-threat-score submissions to the sandbox and how verdicts flow back to the Response & Quarantine Manager and the reputation feedback loop.",
    "3.2 Integration seams": "The analyzer is driven by one public method and two dependency-injected seams. Both seams have a reference implementation that is safe in isolation and a clear production path.",
    "4. Architecture Overview": "This section describes the internal architecture of the Sandbox Analyzer module: its file layout, the eight-stage pipeline, and the component dependency graph.",
    "4.1 Module layout": "The module is organised as a Python package with each pipeline stage implemented in a dedicated file for clear separation of concerns.",
    "4.2 Internal pipeline": "The following diagram traces the eight stages of the sandbox pipeline, including the SAFE and WINDOWS-PE short-circuits that bypass execution.",
    "4.3 Component dependency diagram": "The following diagram shows how SandboxAnalyzer composes the eight stage functions plus the two injected integration seams.",
    "5. Component Design": "This section provides detailed design documentation for every stage, including the routing rules, archive safety layers, kernel-enforced execution limits, the four monitors, the scoring engine, the IOC correlator, the verdict mapper, and the feedback sink.",
    "5.5 Four Monitors": "Each monitor reads a different /proc source at a configurable poll cadence while the sandboxed process runs. The following table lists the data source and the event categories emitted by each monitor.",
    "6. Data Models": "This section defines the typed dataclasses exchanged across the Sandbox Analyzer's public API, including the SandboxReport output and its supporting models.",
    "7. API Specification": "This section documents the public Python API: constructor arguments, primary methods, supporting stage helpers, and an integration example.",
    "8. Input / Output Specification": "This section defines the exact input contract accepted by analyze(), the integration-seam contracts the host must satisfy, and the stable output contract emitted as SandboxReport.to_dict().",
    "9. Sequence Diagrams": "This section provides sequence diagrams for three canonical flows: a SAFE CSV short-circuit, an archive path with double-extension detection, and a script execute-and-monitor flow.",
    "10. Risk Scoring Reference": "The following table lists the default behaviour weights with a rationale per category, the three additive signals bolted onto the behaviour score, and the verdict bands. Weight calibration is an engineering judgement traced in the operational notes.",
    "11. Configuration Reference": "This section enumerates every tunable parameter on SandboxConfig with its default value and meaning, grouped by functional area.",
    "12. Security & Threat Model": "This section enumerates the adversary model with per-vector mitigations, describes the non-repudiation guarantees, specifies the input trust boundary, documents fail-closed configuration validation, and lists the external side effects the sandbox can and cannot have.",
    "13. Performance Characteristics": "This section documents the latency budget per phase, the memory bounds of every internal structure, and the throughput characteristics.",
    "14. Testing Strategy": "This section describes the test-suite layout, how to run it (including the Linux-only integration tier that is auto-skipped on other platforms), and the invariants explicitly enforced.",
    "15. Operations": "This section covers deployment, the operational monitoring checklist, and runbook snippets for common incidents.",
    "16. Risk Assessment": "This section enumerates known risks with their likelihood, impact, and the in-module mitigation.",
}

MS_SECTION_DESCS = {
    "2. Goals and Non-Goals": "This section defines the scope of the Metadata Sanitizer by specifying its intended capabilities and the responsibilities explicitly delegated to other modules in the pipeline.",
    "2.1 Goals": "The following goals define what the Metadata Sanitizer is designed to achieve within the multi-layered detection pipeline.",
    "2.2 Non-Goals": "The following items are explicitly outside the scope of this module and are handled by other components in the system.",
    "3. System Context": "This section describes how the Metadata Sanitizer fits within the overall multi-layered detection pipeline, its position relative to other modules, and the data flow between them.",
    "3.2 Why Separate from Ingestion Interceptor?": "The following table explains why the Metadata Sanitizer is a separate module from the Ingestion Interceptor, despite both dealing with metadata-related operations.",
    "3.3 Data Flow Between Modules": "The following diagram shows the detailed data flow between the Ingestion Interceptor, Game-Theoretic Estimator, Malware Detection Engine, and Metadata Sanitizer, including the specific data fields exchanged.",
    "4. Architecture Overview": "This section describes the internal architecture of the Metadata Sanitizer module, including its file structure, component hierarchy, handler interface design, and how the sanitization rules are organized.",
    "4.1 Module Structure": "The module is organized as a Python package with a handler-based architecture where each file type has a dedicated handler implementing a common interface.",
    "4.2 Component Diagram": "The following diagram shows the internal component structure of the MetadataSanitizer orchestrator, including the handler registry, mode resolver, and the relationship between handlers and sanitization rules.",
    "5. Component Design": "This section provides detailed design documentation for each component, including the orchestrator, mode resolver, handler registry, and all five file-type-specific handlers.",
    "6. Data Models": "This section defines the data models used throughout the sanitization pipeline, including the sanitization mode enum, change records, metadata snapshots, and per-artifact and batch result objects.",
    "7. API Specification": "This section documents the public API surface of the Metadata Sanitizer, including the primary sanitization methods, their parameters, and an integration example with the Ingestion Interceptor.",
    "8. Sequence Diagrams": "This section provides sequence diagrams that illustrate the interaction between components during sanitization operations and the threat-score-driven mode selection logic.",
    "9. Sanitization Rules": "This section defines the declarative sanitization rules for each file type, specifying which metadata fields are stripped, preserved, or flagged in each sanitization mode (strip, selective, audit_only).",
    "10. Configuration Reference": "This section provides the complete reference for all tuneable parameters in the SanitizerConfig dataclass, with defaults and descriptions.",
    "11. Security Considerations": "This section documents the security threat model, defense-in-depth layering, and security properties that the Metadata Sanitizer provides as part of the overall detection pipeline.",
    "13. Dependencies": "This section lists all runtime and optional dependencies, including the graceful degradation behaviour when optional libraries such as Pillow, piexif, pikepdf, or mutagen are absent.",
    "14. Testing Strategy": "This section describes the testing strategy, including test coverage across all components, instructions for running the test suite, and the demo runner for manual verification.",
    "15. Deployment and Operations": "This section covers the deployment topology, monitoring metrics, alert thresholds, and operational procedures for the Metadata Sanitizer in field deployments.",
    "16. Risk Assessment": "This section identifies key risks associated with the Metadata Sanitizer operation and their planned mitigations.",
}


def add_word_toc(doc):
    """Insert a Word TOC field that auto-generates a tabular Table of Contents with page numbers."""
    from docx.oxml import OxmlElement

    h = doc.add_heading("Table of Contents", level=2)

    paragraph = doc.add_paragraph()
    # Begin field
    run1 = paragraph.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    run1._r.append(fld_begin)

    # Field instruction
    run2 = paragraph.add_run()
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = r' TOC \o "1-3" \h \z \u '
    run2._r.append(instr)

    # Separator
    run3 = paragraph.add_run()
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    run3._r.append(fld_sep)

    # Placeholder text
    run4 = paragraph.add_run("Update this field to see table of contents with page numbers (right-click > Update Field)")
    run4.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    run4.font.italic = True
    run4.font.size = Pt(10)

    # End field
    run5 = paragraph.add_run()
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run5._r.append(fld_end)

    doc.add_paragraph()  # spacing


def convert_md_to_docx(md_path, docx_path, diagram_map, prefix, rendered, section_descs=None):
    """Convert a markdown file to a DOCX document."""
    if section_descs is None:
        section_descs = {}
    with open(md_path, "r") as f:
        content = f.read()

    doc = Document()
    setup_styles(doc)
    lines = content.split("\n")
    toc_inserted = False
    skip_toc_items = False

    i = 0
    while i < len(lines):
        line = lines[i]

        # --- HEADING ---
        heading_match = re.match(r'^(#{1,4})\s+(.+)$', line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()

            # Handle Table of Contents: replace markdown TOC with Word TOC field
            if "Table of Contents" in text:
                if not toc_inserted:
                    add_word_toc(doc)
                    toc_inserted = True
                skip_toc_items = True
                i += 1
                continue

            # Stop skipping TOC items when we hit the next real section
            if skip_toc_items and level <= 2 and "Table of Contents" not in text:
                skip_toc_items = False

            h = doc.add_heading(level=level)
            add_formatted_run(h, text)
            i += 1

            # Inject section description if available
            for desc_key, desc_text in section_descs.items():
                if desc_key in text:
                    p = doc.add_paragraph()
                    run = p.add_run(desc_text)
                    run.font.size = Pt(11)
                    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                    p.paragraph_format.space_after = Pt(8)
                    break

            continue

        # --- HORIZONTAL RULE ---
        if re.match(r'^---+\s*$', line.strip()):
            # Add a thin horizontal line
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                '  <w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/>'
                '</w:pBdr>'
            )
            pPr.append(pBdr)
            i += 1
            continue

        # --- CODE BLOCK ---
        if line.strip().startswith("```"):
            lang = line.strip()[3:].strip()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1  # skip closing ```

            code_text = "\n".join(code_lines)

            # Check if this is an ASCII art block
            if is_ascii_art(code_text) and lang == "":
                # Try to match to a Mermaid diagram
                match = find_matching_diagram(code_text, diagram_map, prefix)
                if match:
                    match_str, mermaid_code, width = match
                    key = (prefix, match_str)
                    if key in rendered:
                        png_path, w = rendered[key]
                        add_image_to_doc(doc, png_path, w)
                        continue

                # Try to parse as a box-drawing table
                title, header, rows, footer = parse_box_table(code_text)
                if header and rows:
                    if title:
                        p = doc.add_paragraph()
                        run = p.add_run(title)
                        run.bold = True
                        run.font.size = Pt(12)
                    add_md_table(doc, header, rows)
                    if footer:
                        for fl in footer:
                            p = doc.add_paragraph()
                            add_formatted_run(p, fl)
                            p.paragraph_format.space_before = Pt(2)
                    continue

                # Fallback: keep as code block
                add_code_block(doc, code_text, lang)
            else:
                add_code_block(doc, code_text, lang)
            continue

        # --- TABLE ---
        if line.strip().startswith("|") and "|" in line.strip()[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            header, rows = parse_table_lines(table_lines)
            if header:
                add_md_table(doc, header, rows if rows else [])
            continue

        # --- BLOCKQUOTE ---
        if line.startswith("> ") or line.strip() == ">":
            bq_lines = []
            while i < len(lines) and (lines[i].startswith("> ") or lines[i].strip() == ">" or lines[i].startswith(">")):
                bq_lines.append(lines[i])
                i += 1
            add_blockquote(doc, bq_lines, add_formatted_run, add_md_table, parse_table_lines)
            continue

        # --- UNORDERED LIST ---
        if re.match(r'^(\s*)[-*]\s+', line):
            list_items = []
            while i < len(lines) and re.match(r'^(\s*)[-*]\s+', lines[i]):
                indent = len(re.match(r'^(\s*)', lines[i]).group(1))
                text = re.sub(r'^\s*[-*]\s+', '', lines[i])
                list_items.append((indent, text))
                i += 1
            for indent, text in list_items:
                level = min(indent // 2, 3)
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.left_indent = Inches(0.25 + level * 0.25)
                add_formatted_run(p, text)
            continue

        # --- ORDERED LIST ---
        if re.match(r'^(\s*)\d+\.\s+', line):
            list_items = []
            while i < len(lines) and re.match(r'^(\s*)\d+\.\s+', lines[i]):
                indent = len(re.match(r'^(\s*)', lines[i]).group(1))
                text = re.sub(r'^\s*\d+\.\s+', '', lines[i])
                list_items.append((indent, text))
                i += 1
            # Skip if these are TOC items (contain [link](#anchor) patterns)
            if skip_toc_items:
                continue
            for indent, text in list_items:
                p = doc.add_paragraph(style="List Number")
                add_formatted_run(p, text)
            continue

        # --- BLANK LINE ---
        if line.strip() == "":
            i += 1
            continue

        # --- PARAGRAPH (may span multiple lines) ---
        para_lines = []
        while i < len(lines):
            l = lines[i]
            if l.strip() == "":
                break
            if re.match(r'^#{1,4}\s+', l):
                break
            if l.strip().startswith("```"):
                break
            if l.strip().startswith("|") and "|" in l.strip()[1:]:
                break
            if re.match(r'^---+\s*$', l.strip()):
                break
            if re.match(r'^(\s*)[-*]\s+', l):
                break
            if re.match(r'^\d+\.\s+', l):
                break
            para_lines.append(l)
            i += 1

        if para_lines:
            text = " ".join(para_lines)
            p = doc.add_paragraph()
            add_formatted_run(p, text)

    doc.save(docx_path)
    print(f"  Saved: {docx_path}")


# ===================================================================
# MAIN
# ===================================================================

def main():
    rendered = render_all_diagrams()

    print("\nBuilding Ingestion Interceptor DOCX...")
    convert_md_to_docx(
        os.path.join(DOCS_DIR, "design_ingestion_interceptor.md"),
        os.path.join(DOCS_DIR, "design_ingestion_interceptor.docx"),
        II_DIAGRAM_MAP, "ii", rendered,
        section_descs=II_SECTION_DESCS,
    )

    print("\nBuilding Metadata Sanitizer DOCX...")
    convert_md_to_docx(
        os.path.join(DOCS_DIR, "design_metadata_sanitizer.md"),
        os.path.join(DOCS_DIR, "design_metadata_sanitizer.docx"),
        MS_DIAGRAM_MAP, "ms", rendered,
        section_descs=MS_SECTION_DESCS,
    )

    print("\nBuilding Threat Estimator DOCX...")
    convert_md_to_docx(
        os.path.join(DOCS_DIR, "design_threat_estimator.md"),
        os.path.join(DOCS_DIR, "design_threat_estimator.docx"),
        TE_DIAGRAM_MAP, "te", rendered,
        section_descs=TE_SECTION_DESCS,
    )

    print("\nBuilding Sandbox Analyzer DOCX...")
    convert_md_to_docx(
        os.path.join(DOCS_DIR, "design_sandbox_analyzer.md"),
        os.path.join(DOCS_DIR, "design_sandbox_analyzer.docx"),
        SB_DIAGRAM_MAP, "sb", rendered,
        section_descs=SB_SECTION_DESCS,
    )

    print("\nDone! Files created:")
    for f in [
        "design_ingestion_interceptor.docx",
        "design_metadata_sanitizer.docx",
        "design_threat_estimator.docx",
        "design_sandbox_analyzer.docx",
    ]:
        path = os.path.join(DOCS_DIR, f)
        if os.path.exists(path):
            size = os.path.getsize(path)
            print(f"  {f}: {size/1024:.0f} KB")


if __name__ == "__main__":
    main()
